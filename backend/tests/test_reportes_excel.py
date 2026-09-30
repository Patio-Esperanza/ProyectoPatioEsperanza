import io
import openpyxl
import pytest

from app.schemas.reportes import ReporteTipo
from app.services.reportes_excel import generar_excel_reporte


def test_generar_excel_contenedores_en_patio():
    datos = [
        {
            "contenedor": "MSKU1234567",
            "cliente": "Logística Internacional SA",
            "tipo": "Lleno",
            "tamano": 'CONTENEDOR ESTANDAR 40"',
            "estadia": 5,
            "fecha_entrada": "25/09/2026 10:00:00",
            "sellos": "",
            "patio": "Patio Central",
            "viaje": "ENT123456",
        },
        {
            "contenedor": "TGHU9876543",
            "cliente": "Logística Internacional SA",
            "tipo": "Vacío",
            "tamano": 'CONTENEDOR ESTANDAR 20"',
            "estadia": 2,
            "fecha_entrada": "28/09/2026 15:30:00",
            "sellos": "",
            "patio": "Patio Central",
            "viaje": "ENT654321",
        },
    ]

    buffer = generar_excel_reporte(
        tipo=ReporteTipo.CONTAINERS_IN_YARD,
        datos=datos,
        total_registros=2,
        subtitulo="Snapshot de prueba",
    )

    assert isinstance(buffer, io.BytesIO)
    buffer.seek(0)

    wb = openpyxl.load_workbook(buffer)
    ws = wb.active
    assert ws.title == "Contenedores en Patio"

    # Verificar merge celdas filas 1 y 2
    merged_ranges = [str(r) for r in ws.merged_cells.ranges]
    assert any("A1:I1" in r for r in merged_ranges)
    assert any("A2:I2" in r for r in merged_ranges)

    # Verificar título fila 1
    assert ws["A1"].value == "Contenedores en Patio"
    assert ws["A1"].fill.start_color.rgb == "FF1F4E79" or ws["A1"].fill.start_color.index == "FF1F4E79" or "1F4E79" in str(ws["A1"].fill.start_color.rgb)
    assert ws["A1"].font.name == "Arial"
    assert ws["A1"].font.bold is True

    # Verificar encabezados fila 4
    headers = [ws.cell(row=4, column=col).value for col in range(1, 10)]
    assert "Contenedor" in headers
    assert "Cliente" in headers
    assert "Estadía (Días)" in headers

    # Verificar fila de datos 5 y 6
    assert ws.cell(row=5, column=1).value == "MSKU1234567"
    assert ws.cell(row=6, column=1).value == "TGHU9876543"

    # Verificar alturas
    assert ws.row_dimensions[1].height == 40
    assert ws.row_dimensions[2].height == 22
    assert ws.row_dimensions[4].height == 26
    assert ws.row_dimensions[5].height == 20


def test_generar_excel_todos_los_tipos():
    for tipo in ReporteTipo:
        buffer = generar_excel_reporte(
            tipo=tipo,
            datos=[],
            total_registros=0,
            subtitulo="Prueba vacía",
        )
        assert isinstance(buffer, io.BytesIO)
        buffer.seek(0)
        wb = openpyxl.load_workbook(buffer)
        assert wb.active is not None
