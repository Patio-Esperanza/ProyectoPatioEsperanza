import datetime
import io
from typing import Any
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.config import zona_horaria_app
from app.schemas.reportes import ReporteColumna, ReporteTipo
from app.services.reportes_datos import COLUMNAS_POR_REPORTE, TITULOS_POR_REPORTE


# Paleta de colores estándar para reportes
COLOR_PRIMARIO_HEX = "1F4E79"
COLOR_TEXTO_ENCABEZADO = "FFFFFF"
COLOR_ZEBRA_PAR = "FFFFFF"
COLOR_ZEBRA_IMPAR = "F9F9F9"
COLOR_BORDE = "D9D9D9"


def generar_excel_reporte(
    tipo: ReporteTipo,
    datos: list[dict[str, Any]],
    total_registros: int | None = None,
    subtitulo: str | None = None,
) -> io.BytesIO:
    """
    Genera un archivo Excel .xlsx en memoria con diseño profesional y estructurado
    de acuerdo al catálogo y especificaciones visuales de la empresa.
    """
    columnas: list[ReporteColumna] = COLUMNAS_POR_REPORTE.get(tipo, [])
    titulo_default, subtitulo_default = TITULOS_POR_REPORTE.get(
        tipo, ("Reporte Operativo", "Reporte generado por Patio Esperanza")
    )
    titulo = titulo_default
    sub = subtitulo if subtitulo else subtitulo_default
    num_total = total_registros if total_registros is not None else len(datos)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = titulo[:31]  # Excel limita el nombre de la hoja a 31 caracteres

    # Definición de estilos
    font_titulo = Font(name="Arial", size=16, bold=True, color=COLOR_TEXTO_ENCABEZADO)
    font_subtitulo = Font(name="Arial", size=10, italic=True, color=COLOR_TEXTO_ENCABEZADO)
    font_header = Font(name="Arial", size=11, bold=True, color=COLOR_TEXTO_ENCABEZADO)
    font_data = Font(name="Arial", size=10, color="000000")

    fill_header = PatternFill(
        start_color=COLOR_PRIMARIO_HEX,
        end_color=COLOR_PRIMARIO_HEX,
        fill_type="solid",
    )
    fill_white = PatternFill(
        start_color=COLOR_ZEBRA_PAR,
        end_color=COLOR_ZEBRA_PAR,
        fill_type="solid",
    )
    fill_zebra = PatternFill(
        start_color=COLOR_ZEBRA_IMPAR,
        end_color=COLOR_ZEBRA_IMPAR,
        fill_type="solid",
    )

    side_border = Side(style="thin", color=COLOR_BORDE)
    border_cell = Border(
        left=side_border,
        right=side_border,
        top=side_border,
        bottom=side_border,
    )

    num_columnas = max(len(columnas), 1)
    col_letra_final = get_column_letter(num_columnas)

    # 1. Fila 1: Título Principal
    ws.row_dimensions[1].height = 40
    ws.merge_cells(f"A1:{col_letra_final}1")
    cell_a1 = ws["A1"]
    cell_a1.value = titulo
    cell_a1.font = font_titulo
    cell_a1.alignment = Alignment(horizontal="center", vertical="center")
    cell_a1.fill = fill_header

    for col in range(1, num_columnas + 1):
        c = ws.cell(row=1, column=col)
        c.fill = fill_header

    # 2. Fila 2: Subtítulo y Metadatos
    fecha_emision = datetime.datetime.now(zona_horaria_app()).strftime("%d/%m/%Y %H:%M:%S")
    texto_meta = f"Fecha de Emisión: {fecha_emision} | Total de Registros: {num_total}"
    if sub:
        texto_meta += f" | {sub}"

    ws.row_dimensions[2].height = 22
    ws.merge_cells(f"A2:{col_letra_final}2")
    cell_a2 = ws["A2"]
    cell_a2.value = texto_meta
    cell_a2.font = font_subtitulo
    cell_a2.alignment = Alignment(horizontal="center", vertical="center")
    cell_a2.fill = fill_header

    for col in range(1, num_columnas + 1):
        c = ws.cell(row=2, column=col)
        c.fill = fill_header

    # 3. Fila 3: Espacio en blanco
    ws.row_dimensions[3].height = 12

    # 4. Fila 4: Encabezados de Columnas
    ws.row_dimensions[4].height = 26
    for idx, col in enumerate(columnas, start=1):
        cell = ws.cell(row=4, column=idx)
        cell.value = col.label
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal=col.align or "center", vertical="center")
        cell.border = border_cell

    # 5. Filas 5+: Datos con Zebra Striping
    max_lengths: dict[int, int] = {}
    for idx, col in enumerate(columnas, start=1):
        max_lengths[idx] = len(str(col.label or ""))

    current_row = 5
    for row_idx, item in enumerate(datos):
        ws.row_dimensions[current_row].height = 20
        fill_row = fill_white if row_idx % 2 == 0 else fill_zebra

        for col_idx, col in enumerate(columnas, start=1):
            val = item.get(col.key, "")
            cell = ws.cell(row=current_row, column=col_idx)
            cell.value = val if val is not None else ""
            cell.font = font_data
            cell.fill = fill_row
            cell.alignment = Alignment(horizontal=col.align or "left", vertical="center")
            cell.border = border_cell

            # Calcular ancho máximo
            val_str = str(cell.value)
            if len(val_str) > max_lengths.get(col_idx, 0):
                max_lengths[col_idx] = len(val_str)

        current_row += 1

    # Ajuste de ancho de columnas
    for col_idx in range(1, num_columnas + 1):
        col_letter = get_column_letter(col_idx)
        best_width = max(max_lengths.get(col_idx, 10) + 4, 14)
        ws.column_dimensions[col_letter].width = min(best_width, 45)

    # Guardar en buffer BytesIO
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
